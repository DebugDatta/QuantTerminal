"""Excel (multi-sheet workbook) export for quant-terminal reports.

Function:
    export_excel - Build an openpyxl workbook from a section -> DataFrame
                   mapping with the documented sheet formatting.

Contract sources:
    docs/REPORTS_EXPORT.md       - export_excel(report_data, filename);
                                   multi-sheet workbook built with openpyxl;
                                   sheet inventory and order; bold/frozen
                                   headers; percentage returns / 2-dp prices;
                                   green/red conditional formatting for
                                   returns; auto column widths; confidence
                                   badges rendered as colored fill
    docs/ARCHITECTURE.md         - reports/excel.py pinned as "# export_excel"
    docs/TASK_DIVISION.md        - reports/excel.py -> REPORTS_EXPORT
    docs/FORECASTING_ML.md       - openpyxl listed for Excel export
    docs/MODEL_CONFIDENCE.md     - badge rendered "in Reports ... as a colored
                                   badge in Excel/PDF exports"
    docs/REPORTS_EXPORT.md       - kaleido/Pillow fallback: charts become
                                   tables-only with a warning when the deps
                                   are absent

Notes
-----
- report_data maps a section name to a pandas DataFrame. Sheets are emitted
  in the documented order, skipping sections that are not present; unknown
  dict keys are appended as extra sheets in insertion order.
- Sheet titles are sanitised (max 31 chars, forbidden characters replaced)
  because the documented names are all safe but caller keys may not be.
- The documented currency decoration in headers (e.g. "Close (₹)") has no
  currency input in this signature; a caller that knows the currency may
  pre-decorate its column names. (see REVIEW-LATER)
- Confidence badge rendering: a DataFrame column normalised to
  confidence/badge/level carries per-row values high/medium/low (or the
  emoji forms); the cell(s) are filled with the documented badge colors.
  Badge *computation* (the shared compute_confidence_badge) does not live in
  reports/ (MODEL_CONFIDENCE.md) and is not implemented here.
- Kaleido/Pillow are not present, so chart embedding is not performed; if a
  Charts section is requested the documented tables-only fallback warning is
  emitted.
"""

from __future__ import annotations

import re
import warnings
from datetime import datetime
from pathlib import Path

import pandas as pd

DEFAULT_BASE_DIR = "exports"

# Documented sheet inventory and order (docs/REPORTS_EXPORT.md).
EXCEL_SHEETS = (
    "Executive Summary",
    "Statistics",
    "Returns",
    "Technical Analysis",
    "Risk",
    "Portfolio",
    "Backtest",
    "Forecasting",
    "ML Results",
    "Regimes",
    "Simulation",
    "Factors",
    "Pairs",
    "Charts",
)

# Confidence badge fills (docs/REPORTS_EXPORT.md "Formatting").
BADGE_FILLS = {
    "high": "C6EFCE",
    "medium": "FFEB9C",
    "low": "FFC7CE",
}

RETURNS_NAME_RE = re.compile(r"return|pnl|cagr|yield|drawdown")
PRICE_NAME_RE = re.compile(r"(^|_)(close|open|high|low|price)(_|$)")
VOLUME_NAME_RE = re.compile(r"vol(ume)?(_|$)", re.IGNORECASE)
BADGE_NAME_RE = re.compile(r"^(confidence|badge|level)$")

POSITIVE_FILL = "C6EFCE"
NEGATIVE_FILL = "FFC7CE"

MAX_SHEET_TITLE = 31
FORBIDDEN_SHEET_CHARS = set("[]:*?/\\")
WARN_TABLES_ONLY = (
    "kaleido/Pillow not installed (and no chart-image channel exists in "
    "export_excel); Charts fall back to tables-only."
)


def _timestamp() -> str:
    return datetime.now().replace(microsecond=0).isoformat(sep=" ")


def _load_openpyxl():
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - depends on environment
        raise ImportError(
            "export_excel requires openpyxl (documented dependency for Excel "
            "export in docs/REPORTS_EXPORT.md). Install openpyxl into the "
            "project environment to enable Excel reports."
        ) from exc
    return openpyxl


def _ensure_ext(name: str, ext: str) -> str:
    if not isinstance(name, str) or not name.strip():
        raise ValueError("filename must be a non-empty string")
    stem = name.strip()
    return stem if stem.lower().endswith(ext.lower()) else f"{stem}{ext}"


def _safe_sheet_title(name: str) -> str:
    clean = "".join("_" if c in FORBIDDEN_SHEET_CHARS else c for c in str(name))
    return clean[:MAX_SHEET_TITLE]


def _normalise(col: str) -> str:
    return str(col).strip().lower().replace(" ", "_")


def _badge_level(value) -> str | None:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    s = str(value).strip()
    if s == "🟢":
        return "high"
    if s == "🟡":
        return "medium"
    if s == "🔴":
        return "low"
    low = s.lower()
    if low in {"high", "medium", "low"}:
        return low
    return None


def _write_sheet(ws, df: pd.DataFrame, openpyxl) -> None:
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    include_index = not isinstance(df.index, pd.RangeIndex)
    columns = list(df.columns)
    n_rows, n_cols = df.shape

    header = []
    if include_index:
        header.append(str(df.index.name or ""))
    header.extend([str(c) for c in columns])

    header_font = Font(bold=True)
    for col_idx, label in enumerate(header, start=1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.font = header_font
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 18

    badge_cols = {}
    for j, col in enumerate(columns):
        if BADGE_NAME_RE.match(_normalise(col)):
            badge_cols[j] = _normalise(col)

    ret_cols, price_cols, vol_cols = [], [], []
    for j, col in enumerate(columns):
        n = _normalise(col)
        if n in badge_cols.values():
            continue
        if RETURNS_NAME_RE.search(n):
            ret_cols.append(j)
        elif VOLUME_NAME_RE.search(n):
            vol_cols.append(j)
        elif PRICE_NAME_RE.search(n):
            price_cols.append(j)

    for i in range(n_rows):
        row_idx = i + 2
        if include_index:
            ws.cell(row=row_idx, column=1, value=df.index[i])
        for j, col in enumerate(columns):
            value = df.iloc[i, j]
            cell = ws.cell(row=row_idx, column=1 + (1 if include_index else 0) + j, value=value)
            if j in ret_cols and isinstance(value, (int, float)):
                cell.number_format = "0.00%"
            elif j in price_cols and isinstance(value, (int, float)):
                cell.number_format = "0.00"
            elif j in vol_cols and isinstance(value, (int, float)):
                cell.number_format = "#,##0"
            if j in badge_cols:
                level = _badge_level(value)
                if level:
                    cell.fill = PatternFill(
                        start_color=BADGE_FILLS[level],
                        end_color=BADGE_FILLS[level],
                        fill_type="solid",
                    )

    for j in ret_cols:
        start_col = 1 + (1 if include_index else 0) + j
        col_letter = get_column_letter(start_col)
        rng = f"{col_letter}2:{col_letter}{n_rows + 1}"
        ws.conditional_formatting.add(
            rng,
            CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill(start_color=POSITIVE_FILL, end_color=POSITIVE_FILL, fill_type="solid")),
        )
        ws.conditional_formatting.add(
            rng,
            CellIsRule(operator="lessThan", formula=["0"], fill=PatternFill(start_color=NEGATIVE_FILL, end_color=NEGATIVE_FILL, fill_type="solid")),
        )

    # auto-adjusted column widths (clamped to a readable range)
    for col_idx, label in enumerate(header, start=1):
        col_letter = get_column_letter(col_idx)
        width = len(label)
        for i in range(min(n_rows, 200)):
            v = ws.cell(row=i + 2, column=col_idx).value
            if v is not None:
                width = max(width, len(str(v)))
        ws.column_dimensions[col_letter].width = min(max(width + 2, 8), 50)


def export_excel(
    report_data: dict,
    filename: str,
    base_dir: str | Path = DEFAULT_BASE_DIR,
) -> str:
    """Build and save a multi-sheet Excel report workbook.

    Parameters
    ----------
    report_data : dict[str, pd.DataFrame]
        Section name -> DataFrame of computed analytics. Sheets are written
        in the documented EXCEL_SHEETS order (sections not supplied are
        skipped); extra keys become additional trailing sheets.
    filename : str
        Output filename (.xlsx appended if missing), written under
        base_dir/excel/.
    base_dir : str or Path, default "exports"
        Root that receives the excel/ subfolder.

    Returns
    -------
    str
        Path of the written workbook.
    """
    openpyxl = _load_openpyxl()
    if not isinstance(report_data, dict):
        raise TypeError("report_data must be a dict of section name -> DataFrame")
    if not report_data:
        raise ValueError("report_data must contain at least one section")
    for name, frame in report_data.items():
        if not isinstance(name, str) or not name.strip():
            raise ValueError("report_data keys must be non-empty section names")
        if not isinstance(frame, pd.DataFrame):
            raise TypeError(
                f"report_data values must be pandas DataFrames; got "
                f"{type(frame).__name__} for section {name!r}"
            )

    out_name = _ensure_ext(filename, ".xlsx")
    out_path = Path(base_dir) / "excel" / out_name
    out_path.parent.mkdir(parents=True, exist_ok=True)

    if "Charts" in report_data and not report_data["Charts"].empty:
        warnings.warn(
            "Charts sheet is tables-only: " + WARN_TABLES_ONLY,
            stacklevel=2,
        )

    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    known = set(EXCEL_SHEETS)
    ordered = [s for s in EXCEL_SHEETS if s in report_data]
    ordered += [k for k in report_data if k not in known]
    for section in ordered:
        ws = wb.create_sheet(_safe_sheet_title(section))
        _write_sheet(ws, report_data[section], openpyxl)

    wb.save(out_path)
    return str(out_path)